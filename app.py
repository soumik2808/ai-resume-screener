from flask import Flask, render_template, request, send_file, redirect, url_for
from werkzeug.utils import secure_filename
from sentence_transformers import SentenceTransformer, util
import os
import fitz  # PyMuPDF
import docx
import matplotlib
matplotlib.use('Agg')  # prevent Tkinter thread issues
import matplotlib.pyplot as plt
from io import BytesIO
import base64
from fpdf import FPDF

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

@app.template_filter('b64encode')
def b64encode_filter(data):
    return base64.b64encode(data).decode('utf-8')

model = SentenceTransformer('all-MiniLM-L6-v2')
keywords_list = [
    "python", "flask", "django", "html", "css", "javascript", "sql", "mysql",
    "postgresql", "aws", "azure", "gcp", "docker", "kubernetes", "git", "linux",
    "data analysis", "machine learning", "deep learning", "tensorflow", "pytorch",
    "scikit-learn", "communication", "leadership", "problem solving"
]

def extract_text(file_storage):
    filename = secure_filename(file_storage.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file_storage.save(filepath)

    text = ""
    if filename.endswith('.pdf'):
        with fitz.open(filepath) as doc:
            for page in doc:
                text += page.get_text()
    elif filename.endswith('.docx'):
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            text += para.text + "\n"
    os.remove(filepath)
    return text

def extract_keywords(text, top_k=20):
    embeddings = model.encode([text] + keywords_list, convert_to_tensor=True)
    text_embedding = embeddings[0]
    keyword_embeddings = embeddings[1:]
    cosine_scores = util.cos_sim(text_embedding, keyword_embeddings)[0]
    scored_keywords = list(zip(keywords_list, cosine_scores))
    scored_keywords.sort(key=lambda x: x[1], reverse=True)
    return [kw for kw, _ in scored_keywords[:top_k]]

def generate_pie(matched, total, colors, labels):
    fig, ax = plt.subplots()
    ax.pie([matched, total - matched], labels=labels, colors=colors, autopct='%1.1f%%')
    buf = BytesIO()
    plt.savefig(buf, format='png')
    plt.close(fig)
    return buf.getvalue()

@app.route('/')
def home():
    return render_template("index.html")

@app.route('/screener', methods=['GET', 'POST'])
def screener():
    if request.method == 'POST':
        resume = request.files.get('resume')
        jd = request.files.get('jd')

        resume_text = jd_text = ""
        resume_keywords = jd_keywords = []
        matched_resume = matched_jd = []
        resume_score = 0
        resume_pie = jd_pie = None

        if resume and jd:
            resume_text = extract_text(resume)
            jd_text = extract_text(jd)
            resume_keywords = extract_keywords(resume_text)
            jd_keywords = extract_keywords(jd_text)

            matched_resume = [kw for kw in resume_keywords if kw in jd_keywords]
            matched_jd = [kw for kw in jd_keywords if kw in resume_keywords]
            resume_score = int(len(matched_resume) / len(jd_keywords) * 100) if jd_keywords else 0

            resume_pie = generate_pie(len(matched_resume), len(jd_keywords), ['green', 'red'], ['Matched', 'Missing'])
            jd_pie = generate_pie(len(matched_jd), len(resume_keywords), ['blue', 'orange'], ['Matched', 'Missing'])

        return render_template("screener.html",
            resume_text=resume_text,
            jd_text=jd_text,
            resume_keywords=resume_keywords,
            jd_keywords=jd_keywords,
            matched_resume=matched_resume,
            matched_jd=matched_jd,
            resume_score=resume_score,
            resume_pie=resume_pie,
            jd_pie=jd_pie
        )

    return render_template("screener.html")

@app.route('/builder', methods=['GET', 'POST'])
def builder():
    if request.method == 'POST':
        name = request.form.get('name', '')
        email = request.form.get('email', '')
        phone = request.form.get('phone', '')
        skills = request.form.get('skills', '')
        experience = request.form.get('experience', '')
        education = request.form.get('education', '')

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt=name, ln=True)
        pdf.cell(200, 10, txt=email, ln=True)
        pdf.cell(200, 10, txt=phone, ln=True)
        pdf.multi_cell(0, 10, txt=f"Skills: {skills}")
        pdf.multi_cell(0, 10, txt=f"Experience:\n{experience}")
        pdf.multi_cell(0, 10, txt=f"Education:\n{education}")

        file_path = os.path.join(UPLOAD_FOLDER, "built_resume.pdf")
        pdf.output(file_path)
        response = send_file(file_path, as_attachment=True)
        os.remove(file_path)
        return response

    return render_template("builder.html")

if __name__ == '__main__':
    app.run(debug=True)
